import hashlib
from typing import Any, BinaryIO
from io import BytesIO
from pypdf import PdfReader
from openpyxl import load_workbook

class DocumentParser:
    @staticmethod
    def calculate_sha256(file_stream: BinaryIO) -> str:
        sha256_hash = hashlib.sha256()
        for byte_block in iter(lambda: file_stream.read(4096), b""):
            sha256_hash.update(byte_block)
        file_stream.seek(0)
        return sha256_hash.hexdigest()

class PdfParser:
    @staticmethod
    def parse(file_stream: BinaryIO) -> list[dict[str, Any]]:
        chunks = []
        reader = PdfReader(file_stream)
        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            if text:
                content_hash = hashlib.sha256(text.encode('utf-8')).hexdigest()
                chunks.append({
                    "chunk_index": i,
                    "content": text.strip(),
                    "content_hash": content_hash,
                    "page_number": i + 1,
                })
        file_stream.seek(0)
        return chunks


class ExcelParser:
    ROWS_PER_CHUNK = 5

    @staticmethod
    def parse(file_stream: BinaryIO) -> list[dict[str, Any]]:
        chunks = []

        wb = load_workbook(
            filename=BytesIO(file_stream.read()),
            data_only=True,
            read_only=True,
        )

        chunk_idx = 0

        try:
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]

                # Collect non-empty rows and preserve original row numbers.
                rows = []
                for row_idx, row in enumerate(
                    ws.iter_rows(values_only=True), start=1
                ):
                    values = [
                        "" if cell is None else str(cell).strip()
                        for cell in row
                    ]

                    if not any(values):
                        continue

                    rows.append((row_idx, values))

                if not rows:
                    continue

                # Treat the first non-empty row as the column header.
                header_row_idx, headers = rows[0]
                headers = [
                    value if value else f"Column {i + 1}"
                    for i, value in enumerate(headers)
                ]

                # Group subsequent data rows into context-rich chunks.
                data_rows = rows[1:]

                for start in range(0, len(data_rows), ExcelParser.ROWS_PER_CHUNK):
                    batch = data_rows[
                        start : start + ExcelParser.ROWS_PER_CHUNK
                    ]

                    lines = [
                        f"Worksheet: {sheet_name}",
                        "Columns: " + " | ".join(headers),
                        "",
                    ]

                    for row_idx, values in batch:
                        # Match every value to its corresponding header.
                        cells = []
                        for col_idx, header in enumerate(headers):
                            value = (
                                values[col_idx]
                                if col_idx < len(values)
                                else ""
                            )

                            if value:
                                cells.append(f"{header}: {value}")

                        if cells:
                            lines.append(
                                f"Excel row {row_idx}: " + " | ".join(cells)
                            )

                    content = "\n".join(lines).strip()

                    if not content:
                        continue

                    content_hash = hashlib.sha256(
                        content.encode("utf-8")
                    ).hexdigest()

                    chunks.append(
                        {
                            "chunk_index": chunk_idx,
                            "content": content,
                            "content_hash": content_hash,
                            "sheet_name": sheet_name,
                            "row_start": batch[0][0],
                            "row_end": batch[-1][0],
                        }
                    )

                    chunk_idx += 1
        finally:
            wb.close()
            file_stream.seek(0)

        return chunks